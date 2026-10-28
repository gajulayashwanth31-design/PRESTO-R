from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).parent
OUT = ROOT / "PRESTO-R_Project_Guide_for_Beginners.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(6)
for n in ['Title','Heading 1','Heading 2','Heading 3']:
    styles[n].font.name = 'Aptos Display'
    styles[n].font.color.rgb = RGBColor(0,0,0)
styles['Title'].font.size = Pt(26)
styles['Heading 1'].font.size = Pt(18)
styles['Heading 2'].font.size = Pt(14)
styles['Heading 3'].font.size = Pt(11)

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); pr.append(shd)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); shade(c,'17365D')
        for run in c.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(9)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: shade(cells[i],'EAF0F6')
            for p in cells[i].paragraphs:
                for run in p.runs: run.font.size=Pt(8.5)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph()
    return t
def para(text): doc.add_paragraph(text)
def bullet(text): doc.add_paragraph(text, style='List Bullet')
def heading(text, level=1): doc.add_heading(text, level=level)
def picture(rel, caption, width=6.6):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT/rel), width=Inches(width))
    c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in c.runs: r.italic=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,90,90)

p=doc.add_paragraph(style='Title'); p.add_run('PRESTO-R Project Guide')
p=doc.add_paragraph(); p.add_run('A beginner-friendly, phase-by-phase explanation with tests and actual results').bold=True
para('This guide explains what the project is trying to solve, how its pieces fit together, what was tested, and what the recorded results mean. The core project is a research prototype evaluated in a computer simulation. The main conclusion is mixed: PRESTO-R sharply reduced route switching while keeping delivery close to other learned methods, but it did not meet every predeclared success target. The results do not prove performance on a real production network.')

heading('1. The project in one simple story')
para('Imagine a delivery van choosing one of four roads to reach a shop. A road may become slow, crowded, or blocked. A basic driver waits until the road is already bad, then changes roads. A nervous driver changes roads too often, even when the current road is fine. PRESTO-R tries to look at recent road clues, estimate which network links may become bad soon, and change routes only when the expected benefit is worth the cost of switching.')
para('The project name means Predictive, Stability-Aware Reinforcement Learning Routing. It combines three ideas: prediction (estimate near-future link trouble), reinforcement learning (learn choices from reward), and stability (penalize frequent route changes).')
heading('The gap this project targets',2)
para('The repository frames the gap as a trade-off: reactive routing responds after degradation, while predictive routing can leave early but may flap between routes. PRESTO-R tests whether adding a switching penalty to a predictor-aware learned router can keep delivery quality while reducing unnecessary path changes. This is the project’s stated motivation, not a claim that all existing routers have this weakness.')
heading('What is in scope',2)
for x in ['A discrete-time Python simulator of network traffic and link failures.','One managed source-to-destination flow per episode; other traffic is background load.','Four candidate paths per episode, on NSFNET and a random geometric graph (RGG).','Five synthetic conditions: steady, congestion, node failure, link fluctuation, and mixed events.','A GRU predictor and logistic-regression comparison; heuristic and DQN routing baselines; metrics, statistics, figures, and a report.','Held-out evaluation over 200 scenario seeds, with separate training, validation, and test seed ranges.']: bullet(x)
heading('What is out of scope or not completed',2)
for x in ['Production deployment, real network traffic, or a claim of real-world performance.','Packet-level simulation with NS-3: Phase 7 is explicitly optional and its exit boxes are unchecked.','Multiple managed flows, splitting one flow across paths, or multi-agent coordination.','Real-world failure traces and online learning of the predictor. Failures are synthetic and partly designed with warning precursors.','A full 2×2 experiment isolating prediction and switch penalty: the no-prediction + penalty (reactive_dqn_pen) test cell was not run in the completed grid.','External RQSMR reimplementation. The reactive comparison is an in-project, RQSMR-inspired baseline in a simpler single-flow setup.']: bullet(x)

heading('2. How the system works')
picture('results/figures/F1_architecture.png','Figure 1. Architecture diagram generated by the project.')
table(['Step','In kid-friendly words','Technical meaning'],[
('1. Make a network','Draw computers connected by roads.','Build a graph: nodes are routers; undirected edges are links with capacity and delay.'),
('2. Run traffic','Send one important delivery and some background deliveries.','Fluid simulator advances in 100 ms ticks; background flows create congestion.'),
('3. Observe clues','Read each road’s speed, queue, loss, and whether it is up.','Telemetry is noisy; true link health and future event schedules are hidden from agents.'),
('4. Predict risk','Guess which road may become bad in the next few ticks.','GRU maps a 10-tick, 5-feature history to per-link probability of becoming bad within H ticks.'),
('5. Choose a path','Pick one of four possible routes.','DQN selects a discrete path index from observations built from path features and risk.'),
('6. Score and learn','Good delivery earns a better score; bad service and extra switches cost points.','Reward combines normalized delay, loss, throughput, switch penalty λ, and hop cost.'),
('7. Compare fairly','Try the same situations with every driver.','Deterministic rollouts, paired test seeds, episode metrics, confidence intervals, and corrected tests.')],[.75,2.5,3.35])

heading('Important words, explained')
table(['Term','Plain-language meaning'],[
('Link / node','A link is a connection/road; a node is a router/intersection.'),('Topology','The map of which routers connect to which. NSFNET is a named 14-node, 21-link layout; RGG is a generated layout.'),('Telemetry','Measurements about the network, such as delay, packet loss, queue size, utilization, and up/down state.'),('DQN','A learning algorithm that estimates how good each action is and gradually improves by trial and error.'),('GRU','A neural network designed to use a sequence of past measurements, so it can notice trends.'),('Episode / tick','One simulated trip is an episode; time moves in ticks. Each tick is 100 milliseconds.'),('PDR','Packet/traffic delivery ratio: the part of offered traffic that arrived.'),('Seed','A repeatable random-number starting point. Same seed and actions should recreate the same scenario.'),('λ (lambda)','How costly a route change is in the reward. Higher λ usually encourages fewer switches.')],[1.35,5.25])

heading('3. Project phases and actual checks')
para('The plan numbers phases 0 through 7. Phases 0–6 form the completed core; Phase 7 is optional packet-level validation. A plan can say what a test should do, while a result file says what actually happened. The tables below report recorded status and avoid presenting a planned test as a completed test.')

heading('Phase 0 — Set up a repeatable project',2)
para('Purpose: make sure the code can be installed, configured, and run repeatedly. YAML configuration files hold settings; separate random streams help prevent one random change from altering unrelated parts of the simulation.')
para('Tests/checks: configuration loading and overrides; independent deterministic random streams; dependency installation and quick checks. The report records a fresh Windows environment and 74 tests passing at the Phase 6 release checkpoint. A later Phase 5 note records 97 full-suite tests passing. Treat those as checkpoints at different times, not the same run.')
para('Actual result: Phase 0 is marked complete in the release checklist; README describes pinned dependencies and runnable commands. The original phase-plan checkboxes are planning notes and are not all updated consistently.')

heading('Phase 1 — Build the pretend network',2)
para('Purpose: simulate network maps, traffic, congestion, link/node failures, measurements, route candidates, observations, and reward. The simulator is separated from the reinforcement-learning wrapper so it can be reused.')
para('Tests/checks: topology shape/connectivity and stable link IDs; idle, overloaded, and down-link behavior; precursor and sudden failures; telemetry window order and slopes; state vector shape/range; reward terms; Gymnasium environment checker; fixed-seed repeatability; smoke episode plot.')
para('Actual result: simulator/environment tests exist and Phase 1 is marked complete by later project work. README gives a smoke command that writes results/smoke/episode.png. The final release checklist says a smoke episode was generated. The repository does not provide a single summary of each Phase 1 assertion’s numeric output, so the guide does not invent one.')

heading('Phase 2 — Teach a model to spot link trouble',2)
para('Purpose: from the last W=10 ticks of link measurements, estimate whether the link becomes “bad” in the next H=5 ticks. A link is labelled bad if it is down, its true loss exceeds 0.05, or its true delay exceeds three times its normal delay. Dataset splits are by whole episode to prevent near-identical neighboring ticks leaking across train and test.')
para('Tests/checks: dataset windows and labels do not overlap incorrectly; no scenario seed crosses splits; GRU output shape and probability range; model save/load gives the same output; predictor metrics on held-out data; risk-trace plot.')
para('Actual result: final report shows GRU onset F1 = 0.5296 over all onset samples and PR-AUC = 0.4945. On the precursor-only subgroup, F1 = 0.7476 (1,256 samples); on sudden failures, F1 = 0.6667 (240 samples), with ROC-AUC near chance reported in the phase notes. Logistic regression has overall onset F1 = 0.3949. The planned “onset F1 ≥ 0.7” overall criterion therefore failed. The older Phase 2 plan reports 0.748 for a threshold tuned on the precursor-only subgroup; that is a narrower metric and does not contradict the final overall onset result.')
picture('results/figures/F2_predictor_quality.png','Figure 2. Predictor precision-recall and calibration plot.')

heading('Phase 3 — Add simple drivers and a reactive learner',2)
para('Purpose: create reference strategies and a shared way to replay an episode and compute scores. The shortest-path failover strategy stays on its road unless it goes down. Reactive DQN learns from current conditions but does not get predicted risk.')
para('Tests/checks: heuristic behavior; metric formulas using hand-built examples; same rollout and seed produces identical logs; event logs match generated events; training smoke and validation comparisons. The plan records five reactive DQN training seeds × 300,000 steps.')
para('Actual result: plan notes say reactive DQN exceeded shortest-path failover on validation PDR in node-failure and congestion. Phase 5 held-out NSFNET results show reactive DQN PDR of 0.9662 in node_failure and 0.9721 in congestion; shortest-path failover is 0.8871 and 0.8747 respectively. These final numbers are from the test report, distinct from the earlier validation check.')

heading('Phase 4 — Add prediction and the PRESTO-R stability penalty',2)
para('Purpose: compare a predictive DQN (risk input, λ=0) against PRESTO-R (risk input, λ=0.5). λ was selected on validation data before test evaluation. An oracle-risk agent is a diagnostic upper bound that can see future simulator schedules; it is not a fair deployable competitor.')
para('Tests/checks: validation-only λ sweep; risk permutation diagnostic to check whether the policy uses risk; event-trace comparison; smoke training/config tests; compare learned risk to oracle. The plan says permutation lowered proactive rate from 0.250 to 0.167 on a validation diagnostic, supporting that the policy used risk information.')
para('Actual result: λ=0.5 was selected. Validation notes report 4.30 switches/episode for PRESTO-R versus 21.36 for predictive DQN, with PDR 0.9677 versus 0.9654. On final NSFNET mixed test results, it was 5.197 versus 20.988 switches and PDR 0.9625 versus 0.9608. This is strong evidence for fewer switches at similar delivery in the tested setup.')
picture('results/figures/F5_lambda_tradeoff.png','Figure 3. Test-set trade-off as the switching cost changes.')

heading('Phase 5 — Run the large comparison',2)
para('Purpose: test the strategies on five scenarios, two topologies, held-out seeds, and several settings; compute metrics and statistical comparisons. There are 38,400 episode metric rows in the completed result set.')
para('Tests/checks: resumable evaluation grid; convert raw logs to metrics; summaries by training seed and 95% t-based confidence intervals; paired Wilcoxon signed-rank tests with Holm correction; predefined PASS/FAIL criteria; horizon sweep H=3 and H=10; sudden-failure, 2× noise, and 50% extra-load robustness tests; fresh rerun hash comparison.')
para('Actual result: the reproducibility hash check passed. The final report’s success table says: PRESTO-R vs predictive DQN switching with PDR within the margin passed in mixed and fluctuation; steady delay limit passed; comparison to reactive DQN on failure PDR/recovery failed; overall predictor onset F1 criterion failed. This is a mixed outcome, not a blanket win.')
table(['NSFNET test result','Reactive DQN','Predictive DQN','PRESTO-R','Meaning'],[
('Mixed PDR','0.9619','0.9608','0.9625','All three deliver about 96% of offered traffic.'),
('Mixed switches / episode','18.799','20.988','5.197','PRESTO-R changes route much less.'),
('Node-failure PDR','0.9662','0.9652','0.9667','PRESTO-R is slightly higher by mean; formal success comparison still failed.'),
('Node-failure recovery ticks','8.895','7.646','9.424','PRESTO-R mean recovery is slower than both comparisons here.'),
('Fluctuation PDR','0.9740','0.9726','0.9749','Small differences; stated reactive-comparison target is not fully met.'),
('Steady mean delay (ms)','46.306','46.797','46.277','PRESTO-R is within the 10% limit.')],[1.45,1.05,1.05,1.05,2.0])
para('Interpretation: lower switching is the clearest observed benefit. The delivery ratios are close, but PRESTO-R did not reliably improve recovery versus reactive DQN. Statistical significance and effect size matter; a tiny mean difference does not automatically satisfy a predeclared target.')
picture('results/figures/F4_nsfnet_results.png','Figure 4. Main held-out NSFNET comparisons across scenarios.')
heading('What the extra tests say',3)
para('Horizon sweep: H=3 gave predictor onset F1 0.599; H=10 gave 0.473. In mixed tests, the H=3 agent had PDR 0.9613 and 6.87 switches; H=10 had PDR 0.9616 and 5.56 switches. Longer warning horizon did not improve predictor F1, though route switching stayed low.')
para('Robustness checks: with doubled noise and +50% background load, PRESTO-R still used fewer switches than reactive DQN. Under +50% load, PDR was low for both (0.8736 PRESTO-R, 0.8728 reactive). Sudden failures remove warning signs; these checks do not establish general robustness to real networks.')
para('Statistics: paired tests compare methods on the same scenario seeds. Holm correction reduces false discoveries from testing many comparisons. Across-training-seed standard deviation reflects differences in learned models; it is not the same as variation over test episodes.')

heading('Phase 6 — Make the report and release artifacts',2)
para('Purpose: turn computed data into figures, tables, a report, and model bundle. The plot/report scripts regenerate output from saved results rather than hand-editing numbers.')
para('Tests/checks: fresh environment + quick check; tests; smoke simulation; regenerate figure files; package models and verify SHA-256 checksums; export report PDF; test reproducibility instructions.')
para('Actual result: release checklist records nine figure pairs, a 12-page report PDF, a model bundle of 88 curated files (4 predictor checkpoints and 42 agent policies plus configs), and matching hashes. It also states there is no configured remote release destination, so publishing externally is pending.')

heading('Phase 7 — Optional NS-3 packet validation',2)
para('Purpose: see whether the strategy ranking remains similar in a packet-level simulator. Scope would be smaller: NSFNET, node_failure and mixed, three methods, 30 test episodes per cell. The plan marks all exit criteria unchecked. Therefore NS-3 validation has not been demonstrated as completed and is outside the reported core evidence.')

heading('4. Formulas worth understanding')
heading('4.1 Reward: what the learning agent is trying to maximize',2)
para('The code computes: R = −αD − βL + γT − λS − μE. In the saved standard configuration, α=1, β=4, γ=0, λ=0.5 for PRESTO-R, and μ=0.05. D is normalized delay; L is loss fraction; T is delivered/offered traffic; S is 1 when the route changed this tick and 0 otherwise; E is hops divided by the maximum candidate-path hops. Larger reward is better. Since γ=0 in the configuration, throughput is measured directly in PDR but does not add a separate reward bonus; delivery is still affected through loss/service terms.')
para('Example: changing paths incurs a one-tick penalty of 0.5 in PRESTO-R because λS = 0.5×1. A lost fraction of 0.1 incurs −4×0.1 = −0.4 from the loss term. The exact score also includes delay and hop terms.')
heading('4.2 PDR: how much traffic arrived',2)
para('PDR = total delivered traffic ÷ total offered demand. A PDR of 0.9625 means about 96.25 units arrived for every 100 units offered in aggregate. It is not necessarily a packet count in this fluid simulator.')
heading('4.3 Predictor label',2)
para('y(t,l)=1 if link l becomes bad at any time in (t, t+H], otherwise 0. “Bad” means down, true loss > 0.05, or true delay > 3× base delay. Onset-only evaluation keeps cases where the link is not already bad at t, because already-broken links are easier to classify and are less useful as an early warning.')
heading('4.4 F1 score',2)
para('Precision = true warnings ÷ all warnings. Recall = correctly warned failures ÷ all failures. F1 = 2 × precision × recall ÷ (precision + recall). F1 balances false alarms and missed events. The all-samples predictor result is much higher because already-bad links make the task easier; the onset metric is more relevant to predicting trouble early.')
heading('4.5 Recovery time and proactive rate',2)
para('A proactive response means the selected route has already left the failing link by the failure tick; it counts as a zero-tick recovery in this metric. Otherwise recovery time counts ticks from link failure until a valid selected route has loss ≤ 0.02. Proactive rate = proactive failures avoided ÷ relevant failure events. These definitions are project metrics; they do not mean every real-world outage would recover instantly.')
heading('4.6 Statistical summaries',2)
para('The report first averages episode results within each training seed, then reports mean ± standard deviation across learned-agent seeds. A 95% confidence interval uses the t distribution. Paired Wilcoxon tests compare the same test episodes under two methods; Holm correction adjusts p-values when many comparisons are tested.')

heading('5. How to explain this to faculty')
heading('A short version you can say aloud',2)
para('“My project studies whether a routing agent can anticipate link degradation and avoid excessive route changes. I built a discrete-time network simulator with congestion and synthetic failures, trained a GRU to estimate near-future link risk, and gave that risk to a DQN routing agent. PRESTO-R adds a switching penalty selected on validation data. In held-out simulation tests, PRESTO-R reduced switches substantially—for example, from about 21 to about 5 per episode versus the predictive DQN in NSFNET mixed traffic—while PDR stayed close, around 96%. However, the overall predictor onset F1 target and the predeclared reactive-DQN failure/recovery criteria did not all pass. The work is simulator evidence; packet-level NS-3 and real-network validation remain future work.”')
heading('Likely faculty questions',2)
table(['Question','Clear answer'],[
('What is the main contribution?','Combining predicted per-link risk with a switching-cost reward, then evaluating its delivery/switch trade-off against internal baselines.'),
('What is genuinely new in the result?','In this simulator, the selected penalty brought switching down sharply while keeping PDR close to the predictive agent.'),
('Did the proposed method beat every baseline?','No. It passed the switching/PDR criteria versus predictive DQN in two scenarios and the steady-delay limit, but failed the specified PDR/recovery comparison versus reactive DQN and the overall onset F1 threshold.'),
('Why not claim real-world performance?','The network is a fluid simulator with one managed flow and synthetic events; packet-level NS-3 validation is not completed.'),
('Why include an oracle?','It is a debugging upper bound with access to future simulator conditions, used to see the value of perfect risk information. It is not a deployable competitor.'),
('What would you do next?','Improve onset prediction and test on independent or real traces; complete the missing no-prediction + penalty ablation; then run NS-3 and multi-flow experiments.')],[2.1,4.5])

heading('6. Where to look in the project')
table(['What you need','Repository location'],[
('Project summary and run commands','README.md'),('Phase goals and planned checks','presto-r-plan/phase-0_setup.md through phase-7_ns3_optional.md'),('Final result narrative and tables','results/report.md and results/report.pdf'),('Pass/fail summary','results/success_criteria.md'),('Architecture and result plots','results/figures/'),('Network simulator','presto/sim/'),('Environment, state, reward','presto/env.py, presto/state.py, presto/reward.py'),('Predictor and labels','presto/predictor/ and presto/risk.py'),('Baselines and training','presto/policies/ and presto/agents/'),('Evaluation and statistics','presto/eval/'),('Automated tests','tests/'),('Scenario and agent settings','configs/')],[2.2,4.4])

heading('7. Evidence caveats and honest phrasing')
for x in ['The final report is the source for final held-out results; phase-plan notes sometimes contain earlier validation values or subgroup metrics. State which split and subgroup a number comes from.','The GRU overall onset F1 of 0.5296 is below the 0.7 target; precursor-only F1 0.7476 is a separate, narrower subgroup.','A statistically significant result can still be practically small. Report both the size and the test outcome.','The report’s “38,400 episode rows” refers to aggregated episode metrics across the experiment design, not 38,400 independent network designs or real-world observations.','Use “in this simulator” and “on these tested scenarios” when describing results.']: bullet(x)

para('Primary repository evidence used: README.md; results/report.md; results/success_criteria.md; results/tables; presto-r-plan phase documents; implementation under presto/; automated tests under tests/.')
doc.save(OUT)
print(OUT)
